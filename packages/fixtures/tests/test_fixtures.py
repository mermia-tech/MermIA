import datetime as dt
import json
from collections import Counter
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from mermia_contracts.batch import Grain, NormalizedBatch
from mermia_fixtures import build as b
from mermia_fixtures import restaurant as r
from mermia_fixtures.emit import JSON_DIR, REPO, SR_DIR, orphans, render


def committed(path: Path) -> Any:
    return json.loads((REPO / path).read_text())


def test_committed_files_are_fresh() -> None:
    stale = [
        str(p)
        for p, data in render().items()
        if not (REPO / p).is_file() or (REPO / p).read_bytes() != data
    ]
    assert stale == [], "run: uv run python -m mermia_fixtures"


def test_no_stray_files_in_generated_folders() -> None:
    assert orphans(render()) == [], "run: uv run python -m mermia_fixtures"


def test_sales_lines_are_plausible() -> None:
    # calibration must not leave a negative line, an oversized discount or a one-day spike
    items = b.generate_sales()
    assert all(0 <= x.discount <= x.gross * 0.2 for x in items)
    by_slot: dict[tuple[str, str, int], list[int]] = {}
    for x in items:
        by_slot.setdefault((x.location, x.code, x.day.weekday()), []).append(x.units)
    for slot, units in by_slot.items():
        mean = sum(units) / len(units)
        assert max(units) <= 2 * mean + 3, slot


def test_generation_is_deterministic() -> None:
    assert render() == render()


def test_eval_weekly_figures() -> None:
    # 13 H3 examples 1 and 2
    weekly = committed(JSON_DIR / "expected.json")["weekly_net_sales"]["L1"]
    assert weekly["2026-10-05"] == 18_432_050
    assert weekly["2026-09-28"] == 17_190_500


def test_eval_pastor_cost_and_margin() -> None:
    # 13 H3 example 4: cost $18.40, menu price $65.00; margin on the net price (MR-11 D20):
    # $65.00 less IVA is $56.03, so the margin is 67.2 % and the food cost 32.8 %
    costs = committed(JSON_DIR / "expected.json")["dish_costs_today"]
    pastor = next(c for c in costs if c["code"] == "P001" and c["location"] == "L1")
    assert (
        pastor["cost"],
        pastor["price_gross"],
        pastor["net_price"],
        pastor["margin_pct"],
        pastor["food_cost_pct"],
    ) == (1840, 6500, 5603, "67.2", "32.8")


def test_pastor_meat_costs_less_than_the_eval_proposal() -> None:
    # 13 H3 example 9 proposes $220/kg, which must be a raise over the current price
    assert b.ingredient_price("pastor", r.TODAY) < Decimal(220)


def test_recipes_cover_exactly_the_top_15_items() -> None:
    units: Counter[str] = Counter()
    for x in b.generate_sales():
        units[x.code] += x.units
    top = {code for code, _ in units.most_common(15)}
    assert top == {m.code for m in r.MENU if m.recipe}


def test_dataset_size_matches_the_brief() -> None:
    # 08 C7: 2 locations, 8 weeks, about 40 menu items, about 60 ingredients, 3 suppliers
    data, _, _ = b.build()
    assert len(data["location"]) == 2
    assert (r.LAST_DAY - r.FIRST_DAY).days + 1 == 56
    assert 38 <= len(data["menu_item"]) <= 45
    assert 55 <= len(data["ingredient"]) <= 65
    assert len(data["supplier"]) == 3


def test_cancellation_saturday_stands_out() -> None:
    days = b.daily_totals(b.generate_sales())
    worst = max(days, key=lambda d: d.cancellations_count)
    assert (worst.location, worst.day) == r.CANCEL_SATURDAY
    others = [d.cancellations_count for d in days if d is not worst]
    assert worst.cancellations_count >= 4 * max(others)


def test_ticket_money_adds_up() -> None:
    # ticket grain has the item/daily meaning of net (README Definitions; contract test
    # test_ticket_money_adds_up on ticket_valid.json checks the same rule)
    tickets = committed(JSON_DIR / "sale_ticket.json")
    lines = committed(JSON_DIR / "sale_line.json")
    assert tickets
    for t in tickets:
        own = [x for x in lines if x["ticket_id"] == t["id"]]
        assert own, t["external_id"]
        for x in own:
            assert x["net"] == round(Decimal(x["qty"]) * x["unit_price"]), t["external_id"]
            assert x["tax"] == b.vat_inside(x["net"]), t["external_id"]
        assert t["total"] == sum(x["net"] for x in own), t["external_id"]
        assert t["tax"] == sum(x["tax"] for x in own), t["external_id"]
        assert t["subtotal"] == t["total"] - t["tax"], t["external_id"]


def test_money_is_integer_centavos() -> None:
    for table in ("item_sales_agg", "daily_sales_agg", "menu_item_price"):
        for row in committed(JSON_DIR / f"{table}.json"):
            for col in ("gross", "discount", "net", "net_sales", "tax", "price_gross"):
                if col in row:
                    assert isinstance(row[col], int), (table, col)


GOLDEN = sorted((REPO / SR_DIR).glob("*.golden.json"))


@pytest.mark.parametrize("path", GOLDEN, ids=lambda p: p.name)
def test_golden_is_a_valid_batch(path: Path) -> None:
    batch = NormalizedBatch.model_validate_json(path.read_text())
    rows = batch.item_sales if batch.grain == Grain.item_daily else batch.daily_sales
    assert batch.totals.lines == len(rows)
    if batch.grain == Grain.item_daily:
        assert batch.totals.net_sales == sum(x.net for x in batch.item_sales)
    else:
        assert batch.totals.net_sales == sum(x.net_sales for x in batch.daily_sales)


def test_item_golden_matches_the_seed() -> None:
    batch = committed(SR_DIR / "productos_vendidos_L1_2026-10-11.golden.json")
    loc = b.uid("location", "L1")
    seed = {
        x["external_id"]: x
        for x in committed(JSON_DIR / "item_sales_agg.json")
        if x["location_id"] == loc and x["period_from"] == "2026-10-11" and x["units"] != "0.000000"
    }
    golden = {x["item_external_id"]: x for x in batch["item_sales"]}
    assert golden.keys() == seed.keys()
    for code, g in golden.items():
        s = seed[code]
        assert (g["units"], g["gross"], g["discount"], g["net"]) == (
            s["units"],
            s["gross"],
            s["discount"],
            s["net"],
        ), code


def test_daily_golden_matches_the_seed() -> None:
    batch = committed(SR_DIR / "ventas_por_dia_L1_2026-10-05.golden.json")
    loc = b.uid("location", "L1")
    seed = {
        x["business_date"]: x
        for x in committed(JSON_DIR / "daily_sales_agg.json")
        if x["location_id"] == loc
    }
    fields = (
        "net_sales",
        "tax",
        "tickets",
        "guests",
        "discounts",
        "cancellations_amount",
        "cancellations_count",
        "payment_mix",
        "tips",
    )
    for g in batch["daily_sales"]:
        s = seed[g["business_date"]]
        assert {f: g[f] for f in fields} == {f: s[f] for f in fields}, g["business_date"]
    assert batch["totals"]["net_sales"] == 18_432_050


def table(path: Path, header_row: int = 5) -> list[dict[str, Any]]:
    ws = load_workbook(REPO / path, read_only=True).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    head = [str(c) for c in rows[header_row - 1]]
    return [dict(zip(head, row, strict=True)) for row in rows[header_row:]]


def to_cents(value: Any) -> int:
    return round(Decimal(str(value)) * 100)


def test_item_sample_agrees_with_its_golden() -> None:
    stem = "productos_vendidos_L1_2026-10-11"
    golden = committed(SR_DIR / f"{stem}.golden.json")["item_sales"]
    rows = [x for x in table(SR_DIR / f"{stem}.xlsx") if x["Clave"]]
    assert [x["Clave"] for x in rows] == [g["item_external_id"] for g in golden]
    for x, g in zip(rows, golden, strict=True):
        assert (to_cents(x["Importe"]), to_cents(x["Descuento"]), to_cents(x["Total"])) == (
            g["gross"],
            g["discount"],
            g["net"],
        )


def test_daily_sample_agrees_with_its_golden() -> None:
    stem = "ventas_por_dia_L1_2026-10-05"
    golden = committed(SR_DIR / f"{stem}.golden.json")["daily_sales"]
    rows = table(SR_DIR / f"{stem}.xlsx")
    for x, g in zip(rows, golden, strict=True):
        day = dt.datetime.strptime(x["Fecha"], "%d/%m/%Y").date().isoformat()
        assert day == g["business_date"]
        assert to_cents(x["Venta neta"]) == g["net_sales"]
        assert (
            to_cents(x["Alimentos"]) + to_cents(x["Bebidas"]) - to_cents(x["Descuentos"])
            == g["net_sales"]
        )
